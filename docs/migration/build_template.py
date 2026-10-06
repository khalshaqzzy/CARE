"""Generates Template_Migrasi_Voice_CARE.xlsx for the legacy Voice migration (ADR-0061)."""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Template_Migrasi_Voice_CARE.xlsx")
F = "Arial"
REQ_FILL = PatternFill("solid", fgColor="1F4E78")
OPT_FILL = PatternFill("solid", fgColor="8EA9C1")
EX_FILL = PatternFill("solid", fgColor="F2F2F2")
CARE_FILL = PatternFill("solid", fgColor="C55A11")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

LISTS = {
    "Jenis Voice": ["General", "Private"],
    "Area": ["Karawang 1", "Karawang 2", "Karawang 3", "Sunter 1", "Sunter 2"],
    "Kategori CARE": [
        "Safety", "Environment", "Fasilitas Umum", "Facility Repair",
        "Fasilitas Kerja / Kesulitan Kerja", "Kesejahteraan",
    ],
    "Severity": ["Low", "Medium", "High", "Critical"],
    "Rating": ["1", "2", "3", "4", "5"],
}

# (header, required, width, rule/comment, list name or None, example)
VOICE = [
    ("Platform Asal", True, 16, "Nama platform lama. Tulis sama persis di setiap batch.", None, "Voice App Plant"),
    ("ID Voice Lama", True, 16, "ID unik di platform asal. Jangan diubah antar batch; ID yang sudah pernah dimigrasi akan dilewati.", None, "VC-2025-00123"),
    ("Tanggal Submit", True, 18, "YYYY-MM-DD HH:MM (WIB). Dipakai untuk tren dan filter tanggal di dashboard.", None, "2025-03-14 09:20"),
    ("Jenis Voice", True, 12, "General = ke atasan/department. Private = ke Serikat.", "Jenis Voice", "General"),
    ("NoReg Pelapor", True, 14, "NoReg sesuai data HR, termasuk yang sudah tidak aktif.", None, "1234567"),
    ("Nama Pelapor", True, 22, "Nama saat voice dikirim.", None, "Budi Santoso"),
    ("Directorat Pelapor", False, 22, "Organisasi pelapor saat voice dikirim (dashboard basis Pelapor).", None, "Manufacturing & PE Dir"),
    ("Division Pelapor", True, 22, "Organisasi pelapor saat voice dikirim (dashboard basis Pelapor).", None, "Production Div"),
    ("Department Pelapor", True, 22, "Organisasi pelapor saat voice dikirim (dashboard basis Pelapor).", None, "Welding Dept"),
    ("Section Pelapor", False, 18, "Organisasi pelapor saat voice dikirim (dashboard basis Pelapor).", None, "Welding Section 2"),
    ("Lokasi / Plant (platform lama)", True, 18, "Isi apa adanya dari platform lama.", None, "KRW Plant 1"),
    ("Area CARE", False, 14, "Pilih jika sudah tahu padanannya. Kalau ragu, kosongkan.", "Area", "Karawang 1"),
    ("Detail Lokasi", True, 26, "Maks. 200 karakter.", None, "Gedung A, Line Welding 2"),
    ("Judul", True, 30, "Maks. 150 karakter.", None, "Lantai area welding licin"),
    ("Isi Voice", True, 45, "Maks. 5.000 karakter.", None, "Lantai di area welding line 2 sering licin karena oli, beberapa member hampir terpeleset."),
    ("Kategori (platform lama)", True, 20, "Isi apa adanya dari platform lama. Untuk Private boleh kosong.", None, "K3"),
    ("Kategori CARE", False, 24, "Pilih jika sudah tahu padanannya. Kalau ragu, kosongkan.", "Kategori CARE", "Safety"),
    ("Prioritas (platform lama)", False, 14, "Isi apa adanya jika platform lama punya prioritas/urgensi.", None, "Tinggi"),
    ("Severity (diisi tim CARE)", "care", 14, "Kosongkan. Diisi tim CARE berdasarkan isi voice.", "Severity", "High"),
    ("Directorat Penanganan", False, 22, "Organisasi yang menangani voice (dashboard basis Penanganan, tampilan default).", None, "Manufacturing & PE Dir"),
    ("Division Penanganan", True, 22, "Organisasi yang menangani voice, sesuai nama di CARE. Untuk Private boleh kosong.", None, "Plant Administration Div"),
    ("Department Penanganan", True, 22, "Organisasi yang menangani voice, sesuai nama di CARE. Untuk Private boleh kosong.", None, "Plant GA & SHE Dept"),
    ("Section Penanganan", False, 18, "Isi jika voice ditangani di level section.", None, ""),
    ("NoReg PIC Terakhir", True, 14, "Orang yang terakhir menangani voice ini.", None, "7654321"),
    ("Nama PIC Terakhir", True, 22, "", None, "Andi Wijaya"),
    ("Tanggal Respon Pertama", False, 18, "YYYY-MM-DD HH:MM. Wajib jika voice sudah direspon. Untuk rata-rata waktu respon.", None, "2025-03-14 13:05"),
    ("Tanggal Ditutup", True, 18, "YYYY-MM-DD HH:MM. Untuk rata-rata waktu penyelesaian.", None, "2025-03-20 16:00"),
    ("Catatan Penyelesaian", False, 40, "Maks. 4.000 karakter.", None, "Area sudah dibersihkan dan dipasang anti-slip mat."),
    ("Rating Pelapor (1-5)", False, 12, "Untuk rata-rata skor feedback. Kosongkan jika tidak ada.", "Rating", "5"),
    ("Feedback Pelapor", False, 30, "Maks. 2.000 karakter.", None, "Terima kasih, sudah aman."),
    ("Catatan PIC Platform Lama", False, 30, "Info tambahan yang perlu kami ketahui.", None, ""),
]

LIMITS = {"Detail Lokasi": 200, "Judul": 150, "Isi Voice": 5000,
          "Catatan Penyelesaian": 4000, "Feedback Pelapor": 2000}

wb = Workbook()

# --- Petunjuk ---
ws = wb.active
ws.title = "Petunjuk"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 80
rows = [
    ("Template Migrasi Voice ke CARE", None, "title"),
    ("", None, None),
    ("Cara mengisi", None, "h"),
    ("1", "Isi sheet Voice: satu baris = satu voice yang SUDAH SELESAI (closed). Voice yang masih berjalan dikirim di batch berikutnya setelah selesai.", None),
    ("2", "Baris 2 adalah contoh. Hapus atau timpa saat mengisi.", None),
    ("3", "Kolom \"(platform lama)\" diisi apa adanya. Kolom \"CARE\" boleh dikosongkan, kami petakan bersama.", None),
    ("4", "Tanggal respon dan tanggal ditutup tidak boleh lebih awal dari tanggal submit.", None),
    ("5", "Batch berikutnya cukup berisi voice yang baru selesai. Jika ikut terkirim lagi, voice lama akan dilewati otomatis.", None),
    ("", None, None),
    ("Legenda", None, "h"),
    ("REQ", "Header biru tua = wajib diisi", "req"),
    ("OPT", "Header biru muda = opsional / diisi jika ada", "opt"),
    ("CARE", "Header oranye = dikosongkan, diisi tim CARE", "care"),
    ("EX", "Baris abu-abu = contoh pengisian", "ex"),
    ("", "Klik sel mana pun untuk melihat aturan kolomnya.", None),
    ("", None, None),
    ("Format", None, "h"),
    ("Tanggal", "YYYY-MM-DD HH:MM, zona WIB. Contoh: 2025-03-14 09:20", None),
    ("NoReg", "Tulis sebagai teks, tanpa spasi. Termasuk pelapor yang sudah tidak aktif.", None),
]
for i, (a, b, kind) in enumerate(rows, start=1):
    ca, cb = ws.cell(i, 2, a), ws.cell(i, 3, b)
    ca.font = cb.font = Font(name=F, size=10)
    cb.alignment = Alignment(wrap_text=True, vertical="top")
    if kind == "title":
        ca.font = Font(name=F, size=14, bold=True, color="1F4E78")
    elif kind == "h":
        ca.font = Font(name=F, size=11, bold=True, color="1F4E78")
    elif kind in ("req", "opt", "care", "ex"):
        ca.value = ""
        ca.fill = {"req": REQ_FILL, "opt": OPT_FILL, "care": CARE_FILL, "ex": EX_FILL}[kind]
        ca.border = BORDER


# --- Daftar Pilihan ---
ref = wb.create_sheet("Daftar Pilihan")
ranges = {}
for col, (name, values) in enumerate(LISTS.items(), start=1):
    h = ref.cell(1, col, name)
    h.font = Font(name=F, bold=True, color="FFFFFF")
    h.fill = REQ_FILL
    for r, v in enumerate(values, start=2):
        ref.cell(r, col, v).font = Font(name=F)
    letter = get_column_letter(col)
    ranges[name] = f"='Daftar Pilihan'!${letter}$2:${letter}${len(values) + 1}"
    ref.column_dimensions[letter].width = max(16, max(len(v) for v in values) + 4)
ref.freeze_panes = "A2"


def build(title, cols, rows_hint=2000):
    s = wb.create_sheet(title, index=len(wb.sheetnames) - 1)
    s.freeze_panes = "B2"
    for c, (name, req, width, rule, lst, ex) in enumerate(cols, start=1):
        cell = s.cell(1, c, name + (" *" if req is True else ""))
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill = CARE_FILL if req == "care" else REQ_FILL if req else OPT_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER
        letter = get_column_letter(c)
        s.column_dimensions[letter].width = width
        e = s.cell(2, c, ex)
        e.font = Font(name=F, size=10, italic=True, color="595959")
        e.fill = EX_FILL
        e.border = BORDER
        e.number_format = "@"
        for r in range(3, rows_hint + 2):
            s.cell(r, c).number_format = "@"
        # Column rules show as an input message when a cell is selected. Cell
        # comments are avoided: CARE's importer (ExcelJS) cannot read openpyxl's.
        limit = LIMITS.get(name)
        if lst:
            dv = DataValidation(type="list", formula1=ranges[lst], allow_blank=True,
                                showErrorMessage=True, errorTitle="Pilihan tidak valid",
                                error="Pilih dari daftar.")
        elif limit:
            dv = DataValidation(type="textLength", operator="lessThanOrEqual", formula1=str(limit),
                                allow_blank=True, showErrorMessage=True,
                                errorTitle="Terlalu panjang", error=f"Maksimal {limit} karakter.")
        else:
            dv = DataValidation(allow_blank=True)
        if rule:
            dv.showInputMessage = True
            dv.promptTitle = name[:32]
            dv.prompt = rule[:255]
        s.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{rows_hint + 1}")
    s.row_dimensions[1].height = 42
    s.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"
    return s


build("Voice", VOICE)

wb.move_sheet("Daftar Pilihan", offset=len(wb.sheetnames))
wb.save(OUT)
print(wb.sheetnames)
